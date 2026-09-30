"""Miesięczny raport „Fuel Index”: ORLEN hurt, ARA, kursy walut, ON na stacjach w UE – Excel + HTML + e-mail.

Uruchamianie:
    python report.py                    # raport za poprzedni miesiąc -> reports/fuel_index_RRRR-MM.xlsx i .html
    python report.py --month 2026-08    # wybrany miesiąc
    python report.py --send             # dodatkowo wysyła e-mail (konfiguracja SMTP_* i REPORT_TO, patrz README)
Automatycznie: scheduler.py po udanym odświeżeniu danych wysyła raport za poprzedni miesiąc raz,
od REPORT_DAY dnia miesiąca (domyślnie 4. – wtedy jest już biuletyn KE z ostatniego poniedziałku miesiąca).
"""
import argparse
import io
import os
import smtplib
from datetime import date
from email.message import EmailMessage
from html import escape
from pathlib import Path

import pandas as pd
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill

import db
from sources import ara_manual, orlen
from sources import eu_oil_bulletin as wob

M3_PER_T = 1.1834  # jak w app.py: 1 t ON / 0,845 kg/l
TRANSIT = ["PL", "DE", "CZ", "SK", "LT", "LV", "AT", "HU", "NL", "BE", "LU", "FR", "IT", "ES", "DK", "SE"]
OUT_DIR = Path(__file__).parent / "reports"
MAIL_LOG = "report_mail"  # wpis w fetch_log: message zaczyna się od RRRR-MM wysłanego raportu
NAVY, RED, GREY = "#00417B", "#E2322A", "#5B6B7F"


# ---------------------------------------------------------------- dane
def _s(series: str) -> pd.Series:
    df = db.read_series(series)
    if df.empty:
        return pd.Series(dtype=float, index=pd.DatetimeIndex([]))
    return df.set_index("date")["value"].sort_index()


def _daily_step(s: pd.Series, end: pd.Timestamp) -> pd.Series:
    """Cena obowiązująca w każdym dniu kalendarzowym (ORLEN: notowanie obowiązuje do kolejnej zmiany)."""
    return s.reindex(pd.date_range(s.index.min(), end, freq="D"), method="ffill").dropna()


def _avg(s: pd.Series, p: pd.Period) -> float | None:
    part = s[s.index.to_period("M") == p]
    return float(part.mean()) if not part.empty else None


def _chg(now: float | None, prev: float | None) -> float | None:
    return None if now is None or prev is None or prev == 0 else now / prev - 1


def _mondays(p: pd.Period) -> int:
    return int((pd.date_range(p.start_time, p.end_time.normalize(), freq="D").dayofweek == 0).sum())


def include_ara() -> bool:
    """REPORT_INCLUDE_ARA=0 – bez ARA i premii (licencja OilPriceAPI: notowania ICE tylko do użytku wewnętrznego)."""
    return os.environ.get("REPORT_INCLUDE_ARA", "1").strip() not in ("0", "false", "nie")


def build(month: pd.Period, today: date | None = None, with_ara: bool = True) -> dict:
    """Wszystkie liczby raportu za miesiąc `month`. Brakujące serie -> None (raport powstaje z tego, co jest)."""
    today = today or date.today()
    prev, yago = month - 1, month - 12
    rep = {"month": month, "generated": today, "kpi": [], "notes": []}

    # ORLEN hurt: średnia z dni kalendarzowych
    orl = _s(orlen.SERIES)
    if not orl.empty:
        daily = _daily_step(orl, min(month.end_time.normalize(), pd.Timestamp(today)))
        in_m = orl[orl.index.to_period("M") == month]
        o_now, o_prev, o_y = _avg(daily, month), _avg(daily, prev), _avg(daily, yago)
        d_m = daily[daily.index.to_period("M") == month]
        rep["orlen"] = {
            "avg": o_now, "mm": _chg(o_now, o_prev), "yy": _chg(o_now, o_y),
            "min": float(d_m.min()) if not d_m.empty else None, "max": float(d_m.max()) if not d_m.empty else None,
            "last": float(d_m.iloc[-1]) if not d_m.empty else None, "changes": int(len(in_m)),
            "daily": d_m,
            "hist": pd.Series({p: _avg(daily, p) for p in pd.period_range(month - 12, month, freq="M")}).dropna(),
        }
        rep["kpi"].append(("Ekodiesel ORLEN – hurt (średnia z dni)", o_now, "PLN/m³ netto", _chg(o_now, o_prev),
                           _chg(o_now, o_y), 0))

    # kursy: USD/PLN (Yahoo, dzienne), EUR/PLN (kurs z biuletynu KE)
    usd = _s("usdpln")
    u_now = _avg(usd, month)
    rep["kpi"].append(("USD/PLN (średnia)", u_now, "PLN", _chg(u_now, _avg(usd, prev)), _chg(u_now, _avg(usd, yago)), 4))
    eur = 1 / _s(wob.FX_SERIES)  # biuletyn podaje EUR za 1 PLN
    e_now = _avg(eur, month)
    rep["kpi"].append(("EUR/PLN (średnia z biuletynów KE)", e_now, "PLN", _chg(e_now, _avg(eur, prev)),
                       _chg(e_now, _avg(eur, yago)), 4))

    # ARA i premia (historia ARA jest krótka – liczymy tylko, gdy są notowania w miesiącu)
    ara = _s(ara_manual.SERIES) if with_ara else pd.Series(dtype=float, index=pd.DatetimeIndex([]))
    a_now = _avg(ara, month)
    if a_now is not None:
        a_pln = ara * usd.reindex(ara.index, method="ffill") / M3_PER_T
        ap_now = _avg(a_pln.dropna(), month)
        rep["kpi"].append(("ARA – ICE LS Gasoil (średnia)", a_now, "USD/t", _chg(a_now, _avg(ara, prev)), None, 2))
        if ap_now is not None and not orl.empty:
            prem = (orl.reindex(a_pln.index, method="ffill") - a_pln).dropna()
            rep["kpi"].append(("Premia PL vs ARA (średnia)", _avg(prem, month), "PLN/m³", None, None, 0))
    elif with_ara:
        rep["notes"].append("Brak notowań ARA w tym miesiącu – pominięto ARA i premię.")

    # ON na stacjach w UE (z podatkami, EUR/l)
    eu = db.read_series_like(wob.series_name("brutto", ""))
    if not eu.empty:
        eu["code"] = eu["series"].str[len(wob.series_name("brutto", "")):]
        wide = eu.pivot(index="date", columns="code", values="value").sort_index()
        per = wide.index.to_period("M")
        now_m, prev_m = wide[per == month].mean(), wide[per == prev].mean()
        n_bull = int((per == month).sum())
        codes = [c for c in now_m.dropna().index if c not in wob.AVERAGES and c != "UK"]
        pl = now_m.get("PL")
        table = pd.DataFrame({
            "code": codes,
            "kraj": [wob.COUNTRIES.get(c, c) for c in codes],
            "avg": [now_m[c] for c in codes],
            "mm": [_chg(now_m[c], prev_m.get(c)) for c in codes],
            "vs_pl_1000l": [(now_m[c] - pl) * 1000 if pd.notna(pl) else None for c in codes],
        }).sort_values("avg").reset_index(drop=True)
        table["transit"] = table["code"].isin(TRANSIT)
        rep["eu"] = {"table": table, "bulletins": n_bull, "expected": _mondays(month),
                     "eu_avg": now_m.get("EU"), "pl": pl}
        if pd.notna(pl):
            eur_pl = eur.reindex(wide.index[per == month]).mean()
            rep["kpi"].append(("ON na stacjach – Polska, z podatkami", float(pl), "EUR/l", _chg(pl, prev_m.get("PL")),
                               None, 3))
            if pd.notna(eur_pl):
                rep["kpi"].append(("ON na stacjach – Polska, z podatkami", float(pl * eur_pl), "PLN/l", None, None, 2))
        if pd.notna(now_m.get("EU")):
            rep["kpi"].append(("ON na stacjach – średnia UE-27", float(now_m["EU"]), "EUR/l",
                               _chg(now_m["EU"], prev_m.get("EU")), None, 3))
        if n_bull < rep["eu"]["expected"]:
            rep["notes"].append(f"Biuletyny KE w miesiącu: {n_bull} z {rep['eu']['expected']} poniedziałków – "
                                "średnie UE z niepełnych danych (brakujące tygodnie KE jeszcze nie opublikowała).")
    else:
        rep["notes"].append("Brak danych biuletynu KE w bazie.")

    if month >= pd.Period(today, "M"):
        rep["notes"].append("Miesiąc jeszcze trwa – wartości wstępne.")
    rep["highlights"] = _highlights(rep)
    return rep


def _pct(x: float | None, dec: int = 1) -> str:
    return "—" if x is None or pd.isna(x) else f"{x * 100:+.{dec}f}%".replace(".", ",").replace("-", "−")


def _n(x: float | None, dec: int = 0) -> str:
    return "—" if x is None or pd.isna(x) else f"{x:,.{dec}f}".replace(",", " ").replace(".", ",").replace("-", "−")


def _highlights(rep: dict) -> list[str]:
    out = []
    if o := rep.get("orlen"):
        if o["avg"] is not None:
            out.append(f"Hurt ORLEN średnio {_n(o['avg'])} PLN/m³ netto: {_pct(o['mm'])} m/m, {_pct(o['yy'])} r/r "
                       f"(zakres {_n(o['min'])}–{_n(o['max'])}, zmian cennika: {o['changes']}).")
    if e := rep.get("eu"):
        t = e["table"]
        if not t.empty and pd.notna(e["pl"]):
            pos = int(t.index[t["code"] == "PL"][0]) + 1 if (t["code"] == "PL").any() else None
            cheap = t[t["transit"]].head(3)
            out.append(f"Polska: {_n(e['pl'], 3)} EUR/l na stacjach – {pos}. miejsce z {len(t)} krajów UE "
                       f"(1 = najtaniej)" + (f", {_pct(e['pl'] / e['eu_avg'] - 1)} vs średnia UE." if pd.notna(e["eu_avg"]) else "."))
            if not cheap.empty:
                out.append("Najtaniej na korytarzach tranzytowych: " + ", ".join(
                    f"{r.kraj} {_n(r.avg, 3)} EUR/l" for r in cheap.itertuples()) + ".")
            dear = t[t["transit"]].tail(1)
            if not dear.empty:
                r = dear.iloc[0]
                out.append(f"Najdrożej na korytarzach: {r['kraj']} – {_n(r['vs_pl_1000l'])} EUR więcej za 1000 l niż w Polsce.")
    return out


# ---------------------------------------------------------------- Excel
def to_xlsx(rep: dict) -> bytes:
    kpi = pd.DataFrame([{"Wskaźnik": k, "Wartość": None if v is None else round(v, d), "Jednostka": u,
                         "Zmiana m/m %": None if mm is None else round(mm * 100, 2),
                         "Zmiana r/r %": None if yy is None else round(yy * 100, 2)}
                        for k, v, u, mm, yy, d in rep["kpi"]])
    buf = io.BytesIO()
    with pd.ExcelWriter(buf, engine="openpyxl") as xw:
        kpi.to_excel(xw, sheet_name="Podsumowanie", index=False, startrow=3)
        ws = xw.sheets["Podsumowanie"]
        ws["A1"] = f"Fuel Index – raport miesięczny {rep['month'].strftime('%m.%Y')}"
        ws["A1"].font = Font(bold=True, size=14, color=NAVY[1:])
        ws["A2"] = f"Wygenerowano {rep['generated']:%d.%m.%Y}. Źródła: ORLEN, Komisja Europejska (Weekly Oil Bulletin), Yahoo Finance, OilPriceAPI."
        row = 5 + len(kpi) + 1
        for i, line in enumerate(["Najważniejsze:", *rep["highlights"], *(["Uwagi:", *rep["notes"]] if rep["notes"] else [])]):
            ws.cell(row=row + i, column=1, value=line).font = Font(bold=line.endswith(":"))

        if e := rep.get("eu"):
            t = e["table"]
            pd.DataFrame({
                "Kraj": t["kraj"], "Kod": t["code"], "Średnia ON z podatkami [EUR/l]": t["avg"].round(3),
                "Zmiana m/m %": (t["mm"].astype(float) * 100).round(2),
                "vs Polska na 1000 l [EUR]": t["vs_pl_1000l"].astype(float).round(0),
                "Korytarz tranzytowy": t["transit"].map({True: "tak", False: ""}),
            }).to_excel(xw, sheet_name="Kraje UE", index=False)

        if o := rep.get("orlen"):
            h = o["hist"]
            pd.DataFrame({"Miesiąc": [p.strftime("%Y-%m") for p in h.index], "Średnia ORLEN [PLN/m³]": h.round(2).values}
                         ).to_excel(xw, sheet_name="ORLEN 13 mies.", index=False)
            ws_h = xw.sheets["ORLEN 13 mies."]
            ch = LineChart()
            ch.title, ch.y_axis.title, ch.height, ch.width = "Hurt ORLEN – średnia miesięczna", "PLN/m³", 8, 18
            ch.add_data(Reference(ws_h, min_col=2, min_row=1, max_row=len(h) + 1), titles_from_data=True)
            ch.set_categories(Reference(ws_h, min_col=1, min_row=2, max_row=len(h) + 1))
            ch.legend = None
            ws_h.add_chart(ch, "D2")
            d = o["daily"]
            pd.DataFrame({"Data": d.index.strftime("%Y-%m-%d"), "Cena obowiązująca [PLN/m³]": d.values}
                         ).to_excel(xw, sheet_name="ORLEN dziennie", index=False)

        for ws in xw.sheets.values():
            for cell in ws[4 if ws.title == "Podsumowanie" else 1]:
                if cell.value is not None:
                    cell.font, cell.fill = Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor=NAVY[1:])
            first = 4 if ws.title == "Podsumowanie" else 1  # szerokość wg tabeli, bez tytułu i notatek
            last = first + len(kpi) if ws.title == "Podsumowanie" else ws.max_row
            for col in ws.iter_cols(min_row=first, max_row=last):
                width = max(len(str(c.value or "")) for c in col)
                ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))
    return buf.getvalue()


# ---------------------------------------------------------------- HTML (treść e-maila, do druku jako PDF)
def to_html(rep: dict) -> str:
    m = rep["month"].strftime("%m.%Y")
    td = 'style="padding:6px 10px;border-bottom:1px solid #E3E8EF;"'
    tdr = 'style="padding:6px 10px;border-bottom:1px solid #E3E8EF;text-align:right;font-family:Consolas,monospace;"'
    th = f'style="padding:6px 10px;background:{NAVY};color:#fff;text-align:left;font-weight:600;"'

    def color(x):
        return GREY if x is None or pd.isna(x) or x == 0 else (RED if x > 0 else "#1A8F74")

    kpi_rows = "".join(
        f"<tr><td {td}>{escape(k)}</td><td {tdr}><b>{_n(v, d)}</b> {escape(u)}</td>"
        f'<td {tdr}><span style="color:{color(mm)}">{_pct(mm)}</span></td>'
        f'<td {tdr}><span style="color:{color(yy)}">{_pct(yy)}</span></td></tr>'
        for k, v, u, mm, yy, d in rep["kpi"])
    hl = "".join(f'<li style="margin:0 0 6px;">{escape(x)}</li>' for x in rep["highlights"])
    notes = "".join(f'<p style="margin:4px 0;color:{GREY};font-size:12px;">⚠ {escape(x)}</p>' for x in rep["notes"])

    eu_html = ""
    if e := rep.get("eu"):
        t = e["table"][e["table"]["transit"]]
        rows = "".join(
            f'<tr style="{"background:#FDECEA;font-weight:600;" if r.code == "PL" else ""}">'
            f"<td {td}>{escape(r.kraj)}</td><td {tdr}>{_n(r.avg, 3)}</td>"
            f'<td {tdr}><span style="color:{color(r.mm)}">{_pct(r.mm)}</span></td>'
            f'<td {tdr}>{"—" if r.code == "PL" else _n(r.vs_pl_1000l, 0)}</td></tr>'
            for r in t.itertuples())
        eu_html = (f'<h3 style="color:{NAVY};margin:24px 0 8px;">ON na stacjach – korytarze tranzytowe (średnia {m})</h3>'
                   f'<table style="border-collapse:collapse;width:100%;font-size:14px;"><tr><th {th}>Kraj</th>'
                   f"<th {th}>EUR/l z podatkami</th><th {th}>m/m</th><th {th}>vs PL na 1000 l [EUR]</th></tr>{rows}</table>"
                   f'<p style="color:{GREY};font-size:12px;">Pełna lista 27 krajów w załączniku Excel (arkusz „Kraje UE”). '
                   f"Biuletyny KE w miesiącu: {e['bulletins']} z {e['expected']}.</p>")

    hist_html = ""
    if (o := rep.get("orlen")) and not o["hist"].empty:
        h = o["hist"]
        lo, hi = h.min(), h.max()
        bars = "".join(
            f'<td style="vertical-align:bottom;padding:0 2px;text-align:center;font-size:10px;color:{GREY};">'
            f'<div style="background:{RED if p == rep["month"] else NAVY};height:{100 * v / hi:.0f}px;'
            f'width:100%;border-radius:3px 3px 0 0;"></div>{p.strftime("%m.%y")}</td>'
            for p, v in h.items())
        hist_html = (f'<h3 style="color:{NAVY};margin:24px 0 8px;">Hurt ORLEN – średnie miesięczne, 13 mies. '
                     f"({_n(lo)}–{_n(hi)} PLN/m³)</h3>"
                     f'<table style="border-collapse:collapse;width:100%;height:120px;table-layout:fixed;"><tr>{bars}</tr></table>')

    return f"""<!doctype html><html lang="pl"><head><meta charset="utf-8"><title>Fuel Index {m}</title></head>
<body style="margin:0;background:#F4F6F9;font-family:Segoe UI,Arial,sans-serif;color:#1B2B3C;">
<div style="max-width:760px;margin:0 auto;background:#fff;padding:24px 28px;">
<div style="border-left:6px solid {RED};padding-left:12px;margin-bottom:16px;">
<div style="font-size:12px;color:{GREY};letter-spacing:.05em;text-transform:uppercase;">ID Logistics · monitoring cen paliw</div>
<h1 style="margin:2px 0 0;color:{NAVY};font-size:24px;">Fuel Index – {m}</h1></div>
<h3 style="color:{NAVY};margin:16px 0 8px;">Najważniejsze</h3><ul style="padding-left:20px;margin:0;font-size:15px;line-height:1.5;">{hl}</ul>
{notes}
<h3 style="color:{NAVY};margin:24px 0 8px;">Wskaźniki (średnie miesięczne)</h3>
<table style="border-collapse:collapse;width:100%;font-size:14px;"><tr><th {th}>Wskaźnik</th><th {th}>Średnia</th><th {th}>m/m</th><th {th}>r/r</th></tr>{kpi_rows}</table>
{hist_html}{eu_html}
<p style="color:{GREY};font-size:11px;margin-top:24px;">Wygenerowano {rep['generated']:%d.%m.%Y}. Źródła: ORLEN (hurtowe ceny paliw),
Komisja Europejska – Weekly Oil Bulletin, Yahoo Finance, OilPriceAPI (ICE – tylko do użytku wewnętrznego).
Średnia ORLEN liczona z dni kalendarzowych (cena obowiązuje do kolejnej zmiany). Wzrost na czerwono = drożej dla kupującego.</p>
</div></body></html>"""


# ---------------------------------------------------------------- e-mail
def mail_config() -> dict | None:
    """Konfiguracja z env (na Streamlit Cloud: sekrety najwyższego poziomu trafiają do env). None = niekompletna."""
    cfg = {k: os.environ.get(k, "").strip() for k in
           ("SMTP_HOST", "SMTP_PORT", "SMTP_USER", "SMTP_PASSWORD", "REPORT_FROM", "REPORT_TO")}
    if not (cfg["SMTP_HOST"] and cfg["REPORT_TO"]):
        return None
    cfg["SMTP_PORT"] = int(cfg["SMTP_PORT"] or 587)
    cfg["REPORT_FROM"] = cfg["REPORT_FROM"] or cfg["SMTP_USER"]
    return cfg


def send(rep: dict, cfg: dict | None = None) -> str:
    cfg = cfg or mail_config()
    if cfg is None:
        raise RuntimeError("Brak konfiguracji e-mail: ustaw SMTP_HOST i REPORT_TO (patrz README).")
    m = rep["month"].strftime("%Y-%m")
    to = [a.strip() for a in cfg["REPORT_TO"].split(",") if a.strip()]
    msg = EmailMessage()
    msg["Subject"] = f"Fuel Index {rep['month'].strftime('%m.%Y')} – ceny paliw, raport miesięczny"
    msg["From"], msg["To"] = cfg["REPORT_FROM"], ", ".join(to)
    msg.set_content("Raport w wersji HTML i w załączniku Excel.\n\n" + "\n".join(f"- {x}" for x in rep["highlights"]))
    msg.add_alternative(to_html(rep), subtype="html")
    msg.add_attachment(to_xlsx(rep), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename=f"fuel_index_{m}.xlsx")
    if cfg["SMTP_PORT"] == 465:
        server = smtplib.SMTP_SSL(cfg["SMTP_HOST"], 465, timeout=30)
    else:
        server = smtplib.SMTP(cfg["SMTP_HOST"], cfg["SMTP_PORT"], timeout=30)
        server.starttls()
    with server:
        if cfg["SMTP_USER"]:
            server.login(cfg["SMTP_USER"], cfg["SMTP_PASSWORD"])
        server.send_message(msg)
    return f"{m} wysłany, odbiorcy: {len(to)}"


def auto_send_if_due(today: date | None = None) -> str | None:
    """Raport za poprzedni miesiąc – raz, od REPORT_DAY dnia miesiąca. Wołane po udanym odświeżeniu danych."""
    today = today or date.today()
    if mail_config() is None or today.day < int(os.environ.get("REPORT_DAY", "4") or 4):
        return None
    month = pd.Period(today, "M") - 1
    tag = month.strftime("%Y-%m")
    with db.connect() as con:  # wpisy tego raportu: message zaczyna się od RRRR-MM
        rows = con.execute("SELECT ok FROM fetch_log WHERE series=? AND message LIKE ?", (MAIL_LOG, tag + "%")).fetchall()
    if any(ok for (ok,) in rows) or len(rows) >= 3:  # wysłany albo 3 nieudane próby
        return None
    try:
        msg = send(build(month, today, include_ara()))
        db.log_fetch(MAIL_LOG, True, msg)
    except Exception as e:  # noqa: BLE001 – błąd trafia do logu, nie przerywa odświeżania
        msg = f"{tag}: {type(e).__name__}: {e}"
        db.log_fetch(MAIL_LOG, False, msg)
    return msg


if __name__ == "__main__":
    ap = argparse.ArgumentParser(description="Miesięczny raport Fuel Index (Excel + HTML, opcjonalnie e-mail).")
    ap.add_argument("--month", help="RRRR-MM, domyślnie poprzedni miesiąc")
    ap.add_argument("--out", default=str(OUT_DIR), help="katalog wyjściowy (domyślnie reports/)")
    ap.add_argument("--send", action="store_true", help="wyślij e-mail (SMTP_* i REPORT_TO)")
    a = ap.parse_args()
    mon = pd.Period(a.month, "M") if a.month else pd.Period(date.today(), "M") - 1
    report = build(mon, with_ara=include_ara())
    out = Path(a.out)
    out.mkdir(parents=True, exist_ok=True)
    stem = out / f"fuel_index_{mon.strftime('%Y-%m')}"
    stem.with_suffix(".xlsx").write_bytes(to_xlsx(report))
    stem.with_suffix(".html").write_text(to_html(report), encoding="utf-8")
    print(f"Zapisano {stem}.xlsx i {stem}.html")
    for line in report["highlights"] + report["notes"]:
        print(" -", line)
    if a.send:
        print(send(report))
